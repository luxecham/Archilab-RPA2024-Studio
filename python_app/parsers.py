"""Lecture des exports ETABS / SAFE pour Archilab RPA2024 NoteDeCalcul Studio.

Sources acceptees :
  - texte tabule (export "Display Tables" ou copier/coller ETABS/SAFE) ;
  - classeur Excel .xlsx / .xlsm (un tableau par feuille) ;
  - base Access .mdb / .accdb (access_parser si installe, sinon pilote ODBC).

Le module ne contient aucune geometrie de projet : toutes les entites sont lues dans
les tableaux importes. Les alias de colonnes couvrent ETABS 9.7.4, 18, 21 et SAFE.
"""

from __future__ import annotations

import csv
import io
import math
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable, Iterator, Sequence

__all__ = [
    "ParseError",
    "TabularTable",
    "TabularRow",
    "UnitSystem",
    "Point3D",
    "StoryData",
    "SectionDef",
    "FrameElement",
    "AreaElement",
    "StructuralModel",
    "ModalRatio",
    "BaseReaction",
    "StoryShear",
    "Displacement",
    "ForceResult",
    "MassRigidityCenter",
    "SoilPressure",
    "FootingResult",
    "EtabsResults",
    "SafeResults",
    "ProjectConfig",
    "ParsedProject",
    "read_tabular",
    "read_etabs",
    "read_safe",
    "build_model",
    "read_project",
]


class ParseError(ValueError):
    """Erreur de lecture ou de structure d'un fichier ETABS/SAFE."""


# ---------------------------------------------------------------------------
# Tableaux et resolution d'alias
# ---------------------------------------------------------------------------

def _normalize(value: Any) -> str:
    return re.sub(r"[^A-Z0-9%]", "", str(value).upper())


class TabularRow:
    """Ligne d'un tableau avec resolution tolerante des noms de colonnes."""

    __slots__ = ("_table", "_values")

    def __init__(self, table: "TabularTable", values: Sequence[str]):
        self._table = table
        self._values = list(values)

    def raw(self, *columns: str) -> str | None:
        for column in columns:
            idx = self._table.resolve(column)
            if idx is not None and idx < len(self._values):
                text = str(self._values[idx]).strip()
                if text and text.upper() != "N/A":
                    return text
        return None

    def value(self, *columns: str) -> str | None:
        return self.raw(*columns)

    def number(self, *columns: str) -> float | None:
        text = self.raw(*columns)
        return _to_float(text)

    def percent(self, *columns: str) -> float | None:
        text = self.raw(*columns)
        if text is None:
            return None
        text = text.replace(" ", "").replace(",", ".").rstrip("%")
        return _to_float(text)


def _to_float(text: Any) -> float | None:
    if text is None:
        return None
    if isinstance(text, (int, float)):
        return float(text)
    cleaned = (
        str(text)
        .replace("\u00a0", "")
        .replace(" ", "")
        .replace(",", ".")
        .replace("−", "-")
    )
    cleaned = re.sub(r"[^0-9eE+\-.]", "", cleaned)
    if not cleaned or cleaned in {"-", "+", ".", "-.", "+."}:
        return None
    try:
        return float(cleaned)
    except ValueError:
        return None


class TabularTable:
    """Tableau homogene : nom, entetes, lignes."""

    def __init__(self, name: str, headers: Sequence[str], rows: Sequence[Sequence[str]]):
        self.name = str(name).strip() or "Table"
        self.headers = [str(h).strip() for h in headers]
        self.rows = [list(r) for r in rows if any(str(c).strip() for c in r)]
        self._exact: dict[str, int] = {h: i for i, h in enumerate(self.headers)}
        self._normalized: list[tuple[str, int]] = [
            (_normalize(h), i) for i, h in enumerate(self.headers)
        ]

    @property
    def row_count(self) -> int:
        return len(self.rows)

    def resolve(self, alias: str) -> int | None:
        if alias in self._exact:
            return self._exact[alias]
        low = alias.lower()
        for header, idx in self._exact.items():
            if header.lower() == low:
                return idx
        target = _normalize(alias)
        if not target:
            return None
        for normalized, idx in self._normalized:
            if normalized == target:
                return idx
        for normalized, idx in self._normalized:
            if normalized.startswith(target):
                return idx
        for normalized, idx in self._normalized:
            if target in normalized:
                return idx
        return None

    def iter_rows(self) -> Iterator[TabularRow]:
        for values in self.rows:
            yield TabularRow(self, values)

    def __bool__(self) -> bool:
        return bool(self.headers)

    def __repr__(self) -> str:  # pragma: no cover - confort de debug
        return f"TabularTable({self.name!r}, rows={self.row_count}, cols={len(self.headers)})"


def _find_table(tables: Sequence[TabularTable], *aliases: str) -> TabularTable | None:
    targets = [_normalize(a) for a in aliases]
    for target in targets:
        for table in tables:
            if _normalize(table.name) == target:
                return table
    for target in targets:
        for table in tables:
            if _normalize(table.name).startswith(target):
                return table
    for target in targets:
        for table in tables:
            if target in _normalize(table.name):
                return table
    return None


def _find_tables(tables: Sequence[TabularTable], *aliases: str) -> list[TabularTable]:
    targets = [_normalize(a) for a in aliases]
    found: list[TabularTable] = []
    seen: set[str] = set()
    for table in tables:
        name = _normalize(table.name)
        if any(name == t or name.startswith(t) or t in name for t in targets):
            key = name
            if key not in seen:
                seen.add(key)
                found.append(table)
    return found


# ---------------------------------------------------------------------------
# Unites
# ---------------------------------------------------------------------------

@dataclass(frozen=True)
class UnitSystem:
    length: str = "m"
    force: str = "kN"
    stress: str = "MPa"

    _LENGTH = {
        "M": 1.0, "METER": 1.0, "METRE": 1.0, "CM": 0.01, "MM": 0.001,
        "FT": 0.3048, "FOOT": 0.3048, "IN": 0.0254, "INCH": 0.0254,
    }
    _FORCE = {
        "KN": 1.0, "KILONEWTON": 1.0, "N": 0.001, "NEWTON": 0.001,
        "KGF": 0.00980665, "TONF": 9.80665, "TF": 9.80665,
        "LBF": 0.00444822, "KIP": 4.44822, "KIPS": 4.44822,
    }
    _STRESS = {
        "MPA": 1.0, "N/MM2": 1.0, "NMM2": 1.0, "KPA": 0.001, "GPA": 1000.0,
        "KGF/CM2": 0.0980665, "TONF/M2": 0.00980665, "PSI": 0.00689476,
        "KSI": 6.89476, "T/M2": 0.00980665, "KN/M2": 0.001,
    }

    @classmethod
    def from_etabs(cls, text: str | None) -> "UnitSystem":
        if not text:
            return cls()
        raw = str(text).upper().replace(" ", "")
        parts = [p for p in re.split(r"[,;/|]", raw) if p]
        length = force = stress = None
        for part in parts:
            if part in cls._LENGTH:
                length = part
            elif part in cls._FORCE:
                force = part
            elif part in cls._STRESS:
                stress = part
        # Formes "KN-m", "Tonf-m" ou "Kgf-cm".
        if length is None:
            for token in re.findall(r"[A-Z]+", raw):
                if token in cls._LENGTH:
                    length = token
                    break
        if force is None:
            for token in re.findall(r"[A-Z]+", raw):
                if token in cls._FORCE:
                    force = token
                    break
        if stress is None and "C" in raw and length and force:
            stress = {"M": "MPA", "CM": "KGF/CM2", "MM": "MPA"}.get(length, "MPA")
        return cls(length or "m", force or "kN", stress or "MPa")

    def to_meters(self, value: float | None) -> float:
        if value is None:
            return 0.0
        return float(value) * self._LENGTH.get(self.length.upper(), 1.0)

    def to_centimeters(self, value: float | None) -> float:
        return self.to_meters(value) * 100.0

    def to_kn(self, value: float | None) -> float:
        if value is None:
            return 0.0
        return float(value) * self._FORCE.get(self.force.upper(), 1.0)

    def to_knm(self, value: float | None) -> float:
        return self.to_kn(value)

    def to_mpa(self, value: float | None) -> float:
        if value is None:
            return 0.0
        return float(value) * self._STRESS.get(self.stress.upper().replace(" ", ""), 1.0)

    def to_kpa(self, value: float | None) -> float:
        if value is None:
            return 0.0
        # Une pression se convertit par force / longueur², pas comme une contrainte
        # de materiau exprimee en MPa.
        length = self._LENGTH.get(self.length.upper(), 1.0)
        return self.to_kn(value) / (length * length)

    @property
    def label(self) -> str:
        return f"{self.force}, {self.length}, {self.stress}"


def _detect_units(tables: Sequence[TabularTable]) -> UnitSystem:
    table = _find_table(tables, "Program Control", "Units", "Model Units")
    if table:
        for row in table.iter_rows():
            for alias in ("CurrUnits", "Units", "Current Units", "Present Units"):
                text = row.value(alias)
                if text:
                    return UnitSystem.from_etabs(text)
    return UnitSystem()


# ---------------------------------------------------------------------------
# Lecture des fichiers
# ---------------------------------------------------------------------------

def read_tabular(path: str | Path) -> list[TabularTable]:
    """Lit un export texte, Excel ou Access et retourne ses tableaux."""
    file = Path(path)
    if not file.exists():
        raise ParseError(f"Fichier introuvable : {file}")
    suffix = file.suffix.lower()
    if suffix == ".xls":
        raise ParseError(
            "Le format Excel 97-2003 (.xls) n'est pas pris en charge. "
            "Enregistrer le classeur au format .xlsx puis réimporter."
        )
    if suffix in {".xlsx", ".xlsm"}:
        return _read_excel(file)
    if suffix in {".mdb", ".accdb"}:
        return _read_access(file)
    if suffix in {".csv", ".tsv", ".txt"}:
        return _read_text(file)
    # Extension inconnue : tentative texte puis Excel.
    try:
        return _read_text(file)
    except Exception:
        return _read_excel(file)


def _read_excel(file: Path) -> list[TabularTable]:
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - dependance de packaging
        raise ParseError("openpyxl requis pour les fichiers Excel. Installer : pip install openpyxl") from exc
    try:
        workbook = openpyxl.load_workbook(file, data_only=True, read_only=True)
    except Exception as exc:
        raise ParseError(f"Excel illisible ({file.name}) : {exc}") from exc
    tables: list[TabularTable] = []
    for sheet in workbook.worksheets:
        grid: list[list[str]] = []
        for row in sheet.iter_rows(values_only=True):
            cells = ["" if v is None else str(v).strip() for v in row]
            if any(cells):
                grid.append(cells)
        table = _grid_to_table(sheet.title, grid)
        if table:
            tables.append(table)
    workbook.close()
    if not tables:
        raise ParseError(f"Aucune table exploitable dans {file.name}.")
    return tables


def _read_access(file: Path) -> list[TabularTable]:
    errors: list[str] = []
    try:
        from access_parser import AccessParser  # type: ignore
        parser = AccessParser(str(file))
        tables: list[TabularTable] = []
        for name in parser.catalog:
            if str(name).upper().startswith(("MSYS", "~")):
                continue
            try:
                data = parser.parse_table(name)
            except Exception:
                continue
            if not data:
                continue
            headers = list(data.keys())
            width = max((len(v) for v in data.values()), default=0)
            rows: list[list[str]] = []
            for i in range(width):
                rows.append([
                    "" if data[h][i] is None else str(data[h][i]).strip()
                    for h in headers
                ])
            table = TabularTable(name, headers, rows)
            if table.row_count:
                tables.append(table)
        if tables:
            return tables
    except ImportError:
        errors.append("access_parser non installé")
    except Exception as exc:
        errors.append(f"access_parser : {exc}")

    try:
        import pandas as pd  # type: ignore
        import pyodbc  # type: ignore
        connection = pyodbc.connect(
            r"DRIVER={Microsoft Access Driver (*.mdb, *.accdb)};"
            f"DBQ={file};"
        )
        tables = []
        for row in connection.cursor().tables(tableType="TABLE"):
            name = row.table_name
            df = pd.read_sql(f"SELECT * FROM [{name}]", connection)
            tables.append(TabularTable(
                name,
                [str(c) for c in df.columns],
                df.fillna("").astype(str).values.tolist(),
            ))
        connection.close()
        if tables:
            return tables
    except ImportError:
        errors.append("pyodbc/pandas non installés")
    except Exception as exc:
        errors.append(f"pyodbc : {exc}")

    raise ParseError(
        f"Lecture Access impossible pour {file.name}. "
        "Installer access_parser ou le pilote ODBC Microsoft Access. Details : "
        + " | ".join(errors)
    )


def _read_text(file: Path) -> list[TabularTable]:
    raw = file.read_bytes()
    text = _decode(raw)
    if "TABLE:" in text.upper() and not re.search(r"\t", text):
        parsed = _parse_etabs_text_tables(text)
        if parsed:
            return parsed
    blocks: list[list[str]] = []
    current: list[str] = []
    for line in text.splitlines():
        if line.strip():
            current.append(line.rstrip())
        elif current:
            blocks.append(current)
            current = []
    if current:
        blocks.append(current)

    tables: list[TabularTable] = []
    for index, block in enumerate(blocks):
        table = _block_to_table(block, index)
        if table:
            tables.append(table)
    if not tables:
        raise ParseError(f"Aucun tableau delimite exploitable dans {file.name}.")
    return tables


def _decode(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            text = raw.decode(encoding)
            if "\ufffd" not in text:
                return text
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1", errors="replace")


def _parse_etabs_text_tables(text: str) -> list[TabularTable]:
    """Exports ETABS compacts : 'TABLE:  Nom' suivi d'une ligne d'unites puis des entetes."""
    lines = [line.rstrip() for line in text.splitlines()]
    tables: list[TabularTable] = []
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line.upper().startswith("TABLE:"):
            index += 1
            continue
        name = line.split(":", 1)[1].strip()
        index += 1
        while index < len(lines) and not lines[index].strip():
            index += 1
        if index >= len(lines):
            break
        delimiter = _detect_delimiter(lines[index:index + 3])
        headers = _split_line(lines[index], delimiter)
        # Ligne d'unites optionnelle (ex. "kN, m, C").
        if index + 1 < len(lines):
            candidate = _split_line(lines[index + 1], delimiter)
            if len(candidate) == len(headers) and all(
                re.fullmatch(r"[A-Za-z0-9/%^.\- ]*", c) for c in candidate
            ) and any("KN" in c.upper() or "CM" in c.upper() for c in candidate):
                index += 1
        index += 1
        rows: list[list[str]] = []
        while index < len(lines):
            if lines[index].strip().upper().startswith("TABLE:"):
                break
            if not lines[index].strip():
                index += 1
                if rows and index < len(lines) and lines[index].strip().upper().startswith("TABLE:"):
                    break
                continue
            rows.append(_split_line(lines[index], delimiter))
            index += 1
        if headers and rows:
            tables.append(TabularTable(name, headers, rows))
    return tables


def _block_to_table(block: list[str], index: int) -> TabularTable | None:
    if len(block) < 2:
        return None
    delimiter = _detect_delimiter(block[:4])
    first = _split_line(block[0], delimiter)
    second = _split_line(block[1], delimiter)
    has_title = len(first) <= 1 and len(second) >= 2
    name = first[0] if has_title and first[0].strip() else f"Table {index + 1}"
    header_index = 1 if has_title else 0
    headers = _split_line(block[header_index], delimiter)
    if len(headers) < 2:
        return None
    rows = [
        _split_line(line, delimiter)
        for line in block[header_index + 1:]
        if line.strip()
    ]
    if not rows:
        return None
    width = len(headers)
    rows = [r[:width] + [""] * max(0, width - len(r)) for r in rows]
    return TabularTable(name, headers, rows)


def _grid_to_table(name: str, grid: list[list[str]]) -> TabularTable | None:
    if len(grid) < 2:
        return None
    header_index = next(
        (i for i, row in enumerate(grid) if sum(1 for c in row if c) >= 2),
        None,
    )
    if header_index is None:
        return None
    headers = grid[header_index]
    while headers and not headers[-1]:
        headers = headers[:-1]
    width = len(headers)
    if width < 2:
        return None
    rows = []
    for row in grid[header_index + 1:]:
        if not any(row):
            continue
        rows.append([row[i] if i < len(row) else "" for i in range(width)])
    return TabularTable(name, headers, rows) if rows else None


def _detect_delimiter(lines: Sequence[str]) -> str:
    sample = [line for line in lines if line]
    if any("\t" in line for line in sample):
        return "\t"
    for delimiter in (";", ","):
        if sum(line.count(delimiter) >= 2 for line in sample) >= 2:
            return delimiter
    return "|"


def _split_line(line: str, delimiter: str) -> list[str]:
    if delimiter == "|":
        parts = re.split(r"\s{2,}", line.strip())
    else:
        try:
            parts = next(csv.reader([line], delimiter=delimiter))
        except Exception:
            parts = line.split(delimiter)
    return [str(p).strip().strip('"') for p in parts]


# ---------------------------------------------------------------------------
# Modele geometrique
# ---------------------------------------------------------------------------

@dataclass
class Point3D:
    id: str
    x: float
    y: float
    z: float

    def plan_distance(self, other: "Point3D") -> float:
        return math.hypot(other.x - self.x, other.y - self.y)

    def distance(self, other: "Point3D") -> float:
        return math.hypot(self.plan_distance(other), other.z - self.z)


@dataclass
class StoryData:
    name: str
    height: float
    elevation: float
    is_penthouse: bool = False


@dataclass
class SectionDef:
    name: str
    kind: str  # FRAME_RECTANGULAR, FRAME_CIRCULAR, WALL, SLAB
    material: str = ""
    b: float = 0.0  # cm
    h: float = 0.0  # cm
    diameter: float = 0.0  # cm
    thickness: float = 0.0  # cm
    total_thickness: float = 0.0  # cm
    top_thickness: float = 0.0  # cm
    hollow_height: float = 0.0  # cm
    rib_width: float = 0.0  # cm
    rib_spacing: float = 0.0  # cm

    @property
    def area_cm2(self) -> float:
        if self.kind == "FRAME_CIRCULAR" and self.diameter > 0:
            return math.pi * self.diameter ** 2 / 4.0
        return self.b * self.h

    @property
    def label(self) -> str:
        if self.kind == "FRAME_CIRCULAR":
            return f"Ø{self.diameter:.0f} cm"
        if "FRAME" in self.kind:
            return f"{self.b:.0f}x{self.h:.0f} cm"
        if self.kind == "WALL":
            return f"voile {self.thickness:.0f} cm"
        return f"plancher {self.total_thickness or self.thickness:.0f} cm"


@dataclass
class FrameElement:
    id: str
    kind: str  # COLUMN, BEAM, BRACE
    start: str
    end: str
    section_name: str
    story: str = ""
    length: float = 0.0
    clear_length: float = 0.0


@dataclass
class AreaElement:
    id: str
    kind: str  # WALL, SLAB
    nodes: list[str]
    section_name: str
    pier: str = ""
    story: str = ""


@dataclass
class ElementFamily:
    """Famille d'elements partageant la meme section (ou le meme pier pour un voile)."""

    key: str
    kind: str
    section: SectionDef
    element_ids: list[str] = field(default_factory=list)
    stories: list[str] = field(default_factory=list)
    max_length: float = 0.0
    min_length: float = 0.0
    max_clear_span: float = 0.0
    max_clear_height: float = 0.0
    total_height: float = 0.0
    length_lw: float = 0.0
    pier: str = ""

    @property
    def name(self) -> str:
        return self.key

    @property
    def area_cm2(self) -> float:
        return self.section.area_cm2

    @property
    def geometry_label(self) -> str:
        parts = [self.section.label]
        if self.max_clear_span > 0:
            parts.append(f"L={self.max_clear_span:.2f} m")
        if self.max_clear_height > 0:
            parts.append(f"lcl={self.max_clear_height:.2f} m")
        return " ; ".join(parts)


@dataclass
class StructuralModel:
    joint_name: str
    points: dict[str, Point3D] = field(default_factory=dict)
    stories: list[StoryData] = field(default_factory=list)
    sections: dict[str, SectionDef] = field(default_factory=dict)
    frames: list[FrameElement] = field(default_factory=list)
    areas: list[AreaElement] = field(default_factory=list)
    grid_x: list[float] = field(default_factory=list)
    grid_y: list[float] = field(default_factory=list)
    units: UnitSystem = field(default_factory=UnitSystem)
    warnings: list[str] = field(default_factory=list)

    @property
    def columns(self) -> list[FrameElement]:
        return [f for f in self.frames if f.kind == "COLUMN"]

    @property
    def beams(self) -> list[FrameElement]:
        return [f for f in self.frames if f.kind == "BEAM"]

    @property
    def walls(self) -> list[AreaElement]:
        return [a for a in self.areas if a.kind == "WALL"]

    @property
    def slabs(self) -> list[AreaElement]:
        return [a for a in self.areas if a.kind == "SLAB"]

    @property
    def main_stories(self) -> list[StoryData]:
        return [s for s in self.stories if not s.is_penthouse]

    @property
    def num_stories(self) -> int:
        return len(self.main_stories)

    @property
    def total_height(self) -> float:
        if not self.stories:
            return 0.0
        roof = max(s.elevation for s in self.main_stories)
        if roof > 1e-6:
            return roof
        return sum(s.height for s in self.main_stories)

    @property
    def num_bays_x(self) -> int:
        return max(0, len(self.grid_x) - 1)

    @property
    def num_bays_y(self) -> int:
        return max(0, len(self.grid_y) - 1)

    def story_for_elevation(self, z: float, tolerance: float = 0.05) -> StoryData | None:
        if not self.stories:
            return None
        # Le niveau d'un plancher est son elevation superieure.
        candidates = sorted(self.stories, key=lambda s: s.elevation)
        for story in reversed(candidates):
            if z >= story.elevation - tolerance:
                return story
        return candidates[0]

    def column_families(self) -> list[ElementFamily]:
        return _families(self, [f for f in self.frames if f.kind == "COLUMN"], "COLUMN")

    def beam_families(self) -> list[ElementFamily]:
        return _families(self, [f for f in self.frames if f.kind == "BEAM"], "BEAM")

    def wall_families(self) -> list[ElementFamily]:
        groups: dict[str, list[AreaElement]] = {}
        for wall in self.walls:
            groups.setdefault(wall.pier or wall.section_name, []).append(wall)
        result: list[ElementFamily] = []
        for key, walls in sorted(groups.items()):
            section = self.sections.get(walls[0].section_name)
            if section is None:
                section = SectionDef(walls[0].section_name, "WALL", thickness=20.0)
            nodes = {n for w in walls for n in w.nodes if n in self.points}
            coordinates = [self.points[n] for n in nodes]
            length = 0.0
            if coordinates:
                xs = [p.x for p in coordinates]
                ys = [p.y for p in coordinates]
                length = max(max(xs) - min(xs), max(ys) - min(ys))
            stories = sorted({
                w.story or (self.story_for_elevation(
                    max((self.points[n].z for n in w.nodes if n in self.points), default=0.0)
                ).name if self.story_for_elevation(
                    max((self.points[n].z for n in w.nodes if n in self.points), default=0.0)
                ) else "")
                for w in walls
            })
            heights = [s.height for s in self.stories if s.name in stories]
            total_height = sum(heights) if heights else 0.0
            clear = min(heights) if heights else 0.0
            result.append(ElementFamily(
                key=key,
                kind="WALL",
                section=section,
                element_ids=[w.id for w in walls],
                stories=stories,
                max_length=length,
                min_length=length,
                length_lw=length,
                max_clear_height=clear,
                total_height=total_height,
                pier=key,
            ))
        return result

    def slab_families(self) -> list[ElementFamily]:
        groups: dict[str, list[AreaElement]] = {}
        for slab in self.slabs:
            groups.setdefault(slab.section_name, []).append(slab)
        result: list[ElementFamily] = []
        for key, slabs in sorted(groups.items()):
            section = self.sections.get(key)
            if section is None:
                section = SectionDef(key, "SLAB", thickness=20.0, total_thickness=20.0)
            span = _estimated_slab_span(self, slabs)
            result.append(ElementFamily(
                key=key,
                kind="SLAB",
                section=section,
                element_ids=[s.id for s in slabs],
                stories=sorted({s.story for s in slabs if s.story}),
                max_clear_span=span,
            ))
        return result


def _families(model: StructuralModel, elements: list[FrameElement], kind: str) -> list[ElementFamily]:
    groups: dict[str, list[FrameElement]] = {}
    for element in elements:
        groups.setdefault(element.section_name, []).append(element)
    result: list[ElementFamily] = []
    for section_name, members in sorted(groups.items()):
        section = model.sections.get(section_name)
        if section is None:
            section = SectionDef(section_name, "FRAME_RECTANGULAR", b=30.0, h=40.0)
            model.warnings.append(
                f"Section '{section_name}' absente des proprietes : 30x40 cm retenu pour le controle."
            )
        lengths = [m.length for m in members if m.length > 0]
        clear_spans = [m.clear_length for m in members if m.kind == "BEAM" and m.clear_length > 0]
        clear_heights = [m.clear_length for m in members if m.kind == "COLUMN" and m.clear_length > 0]
        result.append(ElementFamily(
            key=section_name,
            kind=kind,
            section=section,
            element_ids=[m.id for m in members],
            stories=sorted({m.story for m in members if m.story}),
            max_length=max(lengths) if lengths else 0.0,
            min_length=min(lengths) if lengths else 0.0,
            max_clear_span=max(clear_spans) if clear_spans else (max(lengths) if lengths else 0.0),
            max_clear_height=max(clear_heights) if clear_heights else (max(lengths) if lengths else 0.0),
        ))
    return result


def _estimated_slab_span(model: StructuralModel, slabs: list[AreaElement]) -> float:
    spans: list[float] = []
    for slab in slabs:
        points = [model.points[n] for n in slab.nodes if n in model.points]
        if len(points) >= 3:
            xs = [p.x for p in points]
            ys = [p.y for p in points]
            spans.append(max(max(xs) - min(xs), max(ys) - min(ys)))
    if spans:
        return max(spans)
    if len(model.grid_x) >= 2 and len(model.grid_y) >= 2:
        dx = max(_diff(model.grid_x))
        dy = max(_diff(model.grid_y))
        return max(dx, dy)
    return 0.0


def _diff(values: list[float]) -> list[float]:
    ordered = sorted(values)
    return [b - a for a, b in zip(ordered, ordered[1:])]


# ---------------------------------------------------------------------------
# Construction du modele depuis les tableaux
# ---------------------------------------------------------------------------

def build_model(tables: Sequence[TabularTable], joint_name: str) -> StructuralModel:
    if not tables:
        raise ParseError("Le fichier ne contient aucune table exploitable.")
    units = _detect_units(tables)
    model = StructuralModel(joint_name=joint_name, units=units)

    _read_points(tables, units, model)
    _read_stories(tables, units, model)
    _read_sections(tables, units, model)
    _read_frames(tables, units, model)
    _read_areas(tables, units, model)
    _derive_grids(model)
    _finalize_geometry(model)
    if not model.frames and not model.areas:
        raise ParseError(
            "Aucune connectivite d'element trouvee. Exporter les tables "
            "'Connectivity - Frame' et 'Connectivity - Area' depuis ETABS."
        )
    return model


def _read_points(tables: Sequence[TabularTable], units: UnitSystem, model: StructuralModel) -> None:
    table = _find_table(tables, "Point Coordinates", "Joint Coordinates", "Point Object Connectivity")
    if table is None:
        model.warnings.append("Table 'Point Coordinates' absente : topologie non reconstructible.")
        return
    for row in table.iter_rows():
        pid = row.value("UniqueName", "Point", "Joint", "Label", "Name")
        x = row.number("X", "GlobalX")
        y = row.number("Y", "GlobalY")
        z = row.number("Z", "GlobalZ")
        if not pid or x is None or y is None or z is None:
            continue
        model.points[pid] = Point3D(pid, units.to_meters(x), units.to_meters(y), units.to_meters(z))
    if not model.points:
        model.warnings.append("Aucune coordonnee de noeud exploitable.")


def _read_stories(tables: Sequence[TabularTable], units: UnitSystem, model: StructuralModel) -> None:
    table = _find_table(tables, "Story Data", "Stories", "Story Definitions", "Story Names")
    stories: list[StoryData] = []
    if table is not None:
        for row in table.iter_rows():
            name = row.value("Story", "StoryName", "Name", "Level")
            if not name or name.lower() == "base":
                continue
            height = row.number("Height", "StoryHeight", "HK")
            elevation = row.number("Elevation", "Elev", "Z")
            if height is None and elevation is None:
                continue
            stories.append(StoryData(
                name=name,
                height=units.to_meters(height) if height is not None else 0.0,
                elevation=units.to_meters(elevation) if elevation is not None else 0.0,
            ))
    if not stories and model.points:
        elevations = sorted({round(p.z, 4) for p in model.points.values()})
        for index, z in enumerate(elevations):
            next_z = elevations[index + 1] if index + 1 < len(elevations) else z
            height = (next_z - z) if next_z > z else 3.0
            stories.append(StoryData(f"Niveau {index + 1}", height, z))
        model.warnings.append("Table 'Story Data' absente : niveaux deduits des altitudes de noeuds.")
    if stories:
        stories.sort(key=lambda s: s.elevation)
        for story in stories:
            if story.height <= 0 and len(stories) > 1:
                index = stories.index(story)
                if index + 1 < len(stories):
                    story.height = stories[index + 1].elevation - story.elevation
                else:
                    story.height = stories[index].elevation - stories[index - 1].elevation
            story.height = story.height if story.height > 0 else 3.0
        model.stories = stories


def _read_sections(tables: Sequence[TabularTable], units: UnitSystem, model: StructuralModel) -> None:
    frame_tables = [
        t for t in tables
        if any(
            key in _normalize(t.name)
            for key in ("FRAMESECTIONPROPERTIE", "FRAMEPROPERTIE", "BEAMPROPERTIE", "COLUMNPROPERTIE")
        ) and "ASSIGN" not in _normalize(t.name)
    ]
    for table in frame_tables:
        for row in table.iter_rows():
            name = row.value("SectionName", "Section", "Name", "PropName")
            if not name:
                continue
            shape = (row.value("Shape", "SectionShape", "Type") or "RECTANGULAR").upper()
            depth = row.number("t3", "Depth", "H", "D", "Diameter")
            width = row.number("t2", "Width", "B", "D", "Diameter")
            circular = any(k in shape for k in ("CIRC", "ROUND")) or "CIRC" in (row.value("Shape") or "").upper()
            if circular:
                diameter = depth if depth is not None else width
                if diameter is None:
                    continue
                model.sections[name] = SectionDef(
                    name, "FRAME_CIRCULAR",
                    material=row.value("Material", "MatProp") or "",
                    diameter=units.to_centimeters(diameter),
                    b=units.to_centimeters(diameter),
                    h=units.to_centimeters(diameter),
                )
            else:
                if depth is None:
                    continue
                model.sections[name] = SectionDef(
                    name, "FRAME_RECTANGULAR",
                    material=row.value("Material", "MatProp") or "",
                    b=units.to_centimeters(width if width is not None else depth),
                    h=units.to_centimeters(depth),
                )

    area_tables = [
        t for t in tables
        if any(
            key in _normalize(t.name)
            for key in ("WALLPROPERTIE", "SLABPROPERTIE", "DECKPROPERTIE", "AREASECTIONPROPERTIE", "SHELLPROPERTIE")
        ) and "ASSIGN" not in _normalize(t.name)
    ]
    for table in area_tables:
        table_name = _normalize(table.name)
        for row in table.iter_rows():
            name = row.value("Name", "SectionName", "Section", "PropName", "Property")
            if not name:
                continue
            area_type = (row.value("Type", "AreaType", "PropType", "ShellType", "SlabType") or "").upper()
            thickness = (
                row.number("Thickness", "Th", "t3", "SlabThickness")
                or row.number("TotalThickness", "TotalDepth", "Depth")
            )
            total = row.number("TotalThickness", "TotalDepth", "Depth") or thickness
            if thickness is None and total is None:
                continue
            kind = "WALL" if "WALL" in table_name or "WALL" in area_type else "SLAB"
            existing = model.sections.get(name)
            if existing and existing.kind == kind:
                continue
            model.sections[name] = SectionDef(
                name, kind,
                material=row.value("Material", "MatProp") or "",
                thickness=units.to_centimeters(thickness) if thickness is not None else 0.0,
                total_thickness=units.to_centimeters(total) if total is not None else 0.0,
                top_thickness=units.to_centimeters(
                    row.number("TopThickness", "SlabThickness", "CompressionThickness", "H0")
                ),
                hollow_height=units.to_centimeters(
                    row.number("HollowBlockHeight", "BlockHeight", "Hb")
                ),
                rib_width=units.to_centimeters(row.number("RibWidth", "StemWidth", "b0")),
                rib_spacing=units.to_centimeters(row.number("RibSpacing", "Spacing", "Rib")),
            )


def _assignment_map(
    tables: Sequence[TabularTable],
    table_aliases: Sequence[str],
    id_aliases: Sequence[str],
    value_aliases: Sequence[str],
) -> dict[str, str]:
    result: dict[str, str] = {}
    for table in _find_tables(tables, *table_aliases):
        for row in table.iter_rows():
            identifier = row.value(*id_aliases)
            value = row.value(*value_aliases)
            if identifier and value:
                result[identifier] = value
    return result


def _read_frames(tables: Sequence[TabularTable], units: UnitSystem, model: StructuralModel) -> None:
    table = _find_table(tables, "Connectivity - Frame", "Frame Connectivity", "Line Connectivities", "Frame Objects")
    if table is None:
        return
    section_assignments = _assignment_map(
        tables,
        ("Frame Section Assignments", "Line Assigns", "Frame Assignments", "Frame Assigns"),
        ("Frame", "UniqueName", "Line", "Label"),
        ("Section", "SectionName", "PropName", "Property", "AnalSect"),
    )
    story_assignments = _assignment_map(
        tables,
        ("Frame Section Assignments", "Line Assigns", "Frame Assignments", "Frame Assigns"),
        ("Frame", "UniqueName", "Line", "Label"),
        ("Story", "StoryName", "Level"),
    )
    for row in table.iter_rows():
        fid = row.value("Frame", "UniqueName", "Line", "Label")
        start = row.value("Joint", "Joint1", "Joint 1", "Point1", "Point 1", "JointI", "i")
        end = row.value("Joint2", "Joint 2", "Point2", "Point 2", "JointJ", "j")
        if not fid or not start or not end:
            continue
        section = section_assignments.get(fid) or row.value("Section", "SectionName", "PropName", "Property")
        if not section:
            model.warnings.append(f"Element '{fid}' ignore : aucune section affectee.")
            continue
        declared = (row.value("Type", "FrameType", "LineType") or "").upper()
        kind = _infer_frame_kind(start, end, model.points)
        if "COLUMN" in declared:
            kind = "COLUMN"
        elif "BEAM" in declared:
            kind = "BEAM"
        elif "BRACE" in declared:
            kind = "BRACE"
        if kind == "BRACE":
            continue
        story = story_assignments.get(fid) or row.value("Story", "StoryName") or ""
        frame = FrameElement(fid, kind, start, end, section, story)
        model.frames.append(frame)


def _infer_frame_kind(start: str, end: str, points: dict[str, Point3D]) -> str:
    p1 = points.get(start)
    p2 = points.get(end)
    if not p1 or not p2:
        return "OTHER"
    dz = abs(p2.z - p1.z)
    plan = p1.plan_distance(p2)
    if dz > 0.30 and plan < 0.10:
        return "COLUMN"
    if dz < 0.10 and plan > 0.10:
        return "BEAM"
    return "BRACE"


def _read_areas(tables: Sequence[TabularTable], units: UnitSystem, model: StructuralModel) -> None:
    table = _find_table(tables, "Connectivity - Area", "Area Connectivity", "Area Connectivities", "Area Objects")
    if table is None:
        return
    section_assignments = _assignment_map(
        tables,
        ("Area Section Assignments", "Area Assigns", "Shell Assignments", "Wall Assignments"),
        ("Area", "UniqueName", "Label", "Wall"),
        ("Section", "SectionName", "PropName", "Property", "ShellType"),
    )
    story_assignments = _assignment_map(
        tables,
        ("Area Section Assignments", "Area Assigns", "Shell Assignments", "Wall Assignments"),
        ("Area", "UniqueName", "Label", "Wall"),
        ("Story", "StoryName", "Level"),
    )
    pier_assignments = _assignment_map(
        tables,
        ("Wall Pier Assignments", "Pier Assignments", "Pier Labels", "Pier Names"),
        ("Area", "UniqueName", "Label", "Wall"),
        ("Pier", "PierName", "PierLabel"),
    )
    for row in table.iter_rows():
        aid = row.value("Area", "UniqueName", "Label")
        if not aid:
            continue
        nodes = [
            row.value(f"Point{i}", f"Point {i}", f"Joint{i}", f"Joint {i}")
            for i in range(1, 9)
        ]
        node_ids = [n for n in nodes if n]
        if len(node_ids) < 3:
            continue
        section = section_assignments.get(aid) or row.value("Section", "SectionName", "PropName", "Property")
        if not section:
            model.warnings.append(f"Surface '{aid}' ignoree : aucune propriete affectee.")
            continue
        declared = (row.value("AreaType", "Area Type", "Type", "ShellType") or "").upper()
        kind = _infer_area_kind(node_ids, model.points)
        if "WALL" in declared:
            kind = "WALL"
        elif any(k in declared for k in ("FLOOR", "ROOF", "SLAB", "DECK")):
            kind = "SLAB"
        story = story_assignments.get(aid) or row.value("Story", "StoryName") or ""
        model.areas.append(AreaElement(
            aid, kind, node_ids, section,
            pier=pier_assignments.get(aid) or row.value("Pier", "PierName", "PierLabel") or "",
            story=story,
        ))


def _infer_area_kind(nodes: Sequence[str], points: dict[str, Point3D]) -> str:
    coords = [points[n] for n in nodes if n in points]
    if len(coords) < 3:
        return "OTHER"
    dz = max(p.z for p in coords) - min(p.z for p in coords)
    dx = max(p.x for p in coords) - min(p.x for p in coords)
    dy = max(p.y for p in coords) - min(p.y for p in coords)
    return "WALL" if dz > 0.30 and min(dx, dy) < 1.0 else "SLAB"


def _derive_grids(model: StructuralModel) -> None:
    table_x: list[float] = []
    table_y: list[float] = []
    if not model.grid_x and model.points:
        xs = sorted({round(p.x, 3) for p in model.points.values()})
        ys = sorted({round(p.y, 3) for p in model.points.values()})
        if len(xs) >= 2:
            model.grid_x = xs
        if len(ys) >= 2:
            model.grid_y = ys
    model.grid_x = _merge_close(model.grid_x)
    model.grid_y = _merge_close(model.grid_y)


def _merge_close(values: Sequence[float], tolerance: float = 0.05) -> list[float]:
    result: list[float] = []
    for value in sorted(values):
        if not result or abs(value - result[-1]) > tolerance:
            result.append(value)
    return result


def _finalize_geometry(model: StructuralModel) -> None:
    for frame in model.frames:
        p1 = model.points.get(frame.start)
        p2 = model.points.get(frame.end)
        if not p1 or not p2:
            continue
        frame.length = p1.distance(p2)
        if not frame.story:
            story = model.story_for_elevation(max(p1.z, p2.z))
            frame.story = story.name if story else ""
        if frame.kind == "COLUMN":
            frame.clear_length = frame.length
        else:
            frame.clear_length = frame.length

    # Deduction du niveau pour les surfaces (altitude du plancher superieur).
    for area in model.areas:
        if area.story:
            continue
        z = max((model.points[n].z for n in area.nodes if n in model.points), default=0.0)
        story = model.story_for_elevation(z)
        area.story = story.name if story else ""

    # Hauteur libre de poteau : hauteur d'etage moins la retombee de poutre.
    beam_depths = {
        name: sec.h for name, sec in model.sections.items() if sec.kind == "FRAME_RECTANGULAR"
    }
    for frame in model.frames:
        if frame.kind != "COLUMN":
            continue
        story = next((s for s in model.stories if s.name == frame.story), None)
        if story is None and frame.story:
            # Le nom de niveau d'un poteau est son niveau superieur.
            story = model.story_for_elevation(max(
                model.points[frame.start].z, model.points[frame.end].z
            ))
        if story is None:
            continue
        depth = 0.0
        for beam in model.beams:
            if beam.story == story.name:
                sec = model.sections.get(beam.section_name)
                if sec:
                    depth = max(depth, sec.h)
        frame.clear_length = max(0.2, story.height - depth / 100.0)

    # Portee libre des poutres : longueur moins demi-largeurs d'appui.
    for frame in model.frames:
        if frame.kind != "BEAM":
            continue
        p1 = model.points.get(frame.start)
        p2 = model.points.get(frame.end)
        if not p1 or not p2:
            continue
        supports = [0.0]
        for column in model.columns:
            if column.story != frame.story:
                continue
            endpoint = model.points.get(column.start)
            if endpoint and endpoint.plan_distance(p1) < 0.15:
                sec = model.sections.get(column.section_name)
                if sec:
                    supports.append(max(sec.b, sec.h) / 100.0)
            endpoint = model.points.get(column.end)
            if endpoint and endpoint.plan_distance(p2) < 0.15:
                sec = model.sections.get(column.section_name)
                if sec:
                    supports.append(max(sec.b, sec.h) / 100.0)
        frame.clear_length = max(0.1, frame.length - sum(supports) / 2.0)


# ---------------------------------------------------------------------------
# Resultats ETABS
# ---------------------------------------------------------------------------

@dataclass
class ModalRatio:
    mode: int
    period: float
    ux: float
    uy: float
    uz: float
    sum_ux: float
    sum_uy: float
    sum_uz: float
    rx: float
    ry: float
    rz: float


@dataclass
class BaseReaction:
    load_case: str
    fx: float
    fy: float
    fz: float
    story: str = ""


@dataclass
class StoryShear:
    story: str
    load_case: str
    location: str
    vx: float
    vy: float
    p: float

    @property
    def is_bottom(self) -> bool:
        return self.location.strip().lower().startswith("bottom")


@dataclass
class Displacement:
    story: str
    load_case: str
    ux: float
    uy: float
    uz: float
    diaphragm: str = ""


@dataclass
class ForceResult:
    element: str
    load_case: str
    story: str
    station: float
    p: float
    v2: float
    v3: float
    m2: float
    m3: float
    length: float = 0.0


@dataclass
class MassRigidityCenter:
    story: str
    diaphragm: str
    xcm: float
    ycm: float
    xcr: float
    ycr: float
    mass_x: float = 0.0
    mass_y: float = 0.0


@dataclass
class SoilPressure:
    label: str
    load_case: str
    pressure: float  # kPa
    story: str = ""
    displacement: float = 0.0  # m
    area: float = 0.0  # m2


@dataclass
class FootingResult:
    label: str
    b: float
    length: float
    thickness: float
    n: float
    mx: float = 0.0
    my: float = 0.0
    q_adm: float = 0.0
    q_max: float = 0.0
    q_min: float = 0.0
    settlement: float = 0.0


@dataclass
class EtabsResults:
    modal: list[ModalRatio] = field(default_factory=list)
    base_reactions: list[BaseReaction] = field(default_factory=list)
    story_shears: list[StoryShear] = field(default_factory=list)
    displacements: list[Displacement] = field(default_factory=list)
    column_forces: list[ForceResult] = field(default_factory=list)
    beam_forces: list[ForceResult] = field(default_factory=list)
    pier_forces: list[ForceResult] = field(default_factory=list)
    centers: list[MassRigidityCenter] = field(default_factory=list)
    soil_pressures: list[SoilPressure] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    @property
    def seismic_weight(self) -> float | None:
        rows = [r for r in self.story_shears if _is_weight_case(r.load_case)]
        if not rows:
            rows = [r for r in self.base_reactions if _is_weight_case(r.load_case)]
        if not rows:
            return None
        if rows and isinstance(rows[0], StoryShear):
            return max(abs(r.p) for r in rows)
        return max(abs(r.fz) for r in rows if isinstance(r, BaseReaction))

    def base_shear(self, direction: str) -> float | None:
        case = "EX" if direction.upper() == "X" else "EY"
        rows = [r for r in self.story_shears if _matches_case(r.load_case, case) and r.is_bottom]
        if not rows:
            rows = [r for r in self.story_shears if _matches_case(r.load_case, case)]
        if rows:
            return max(abs(r.vx if direction.upper() == "X" else r.vy) for r in rows)
        rows = [r for r in self.base_reactions if _matches_case(r.load_case, case)]
        if rows:
            return max(abs(r.fx if direction.upper() == "X" else r.fy) for r in rows)
        return None

    def cumulative_weights(self) -> dict[str, float]:
        rows = [r for r in self.story_shears if _is_weight_case(r.load_case)]
        by_story: dict[str, float] = {}
        for row in rows:
            key = row.story.strip().upper()
            by_story[key] = max(by_story.get(key, 0.0), abs(row.p))
        if by_story:
            return by_story
        # A defaut : repartition uniforme du poids total sur les niveaux.
        weight = self.seismic_weight
        stories = {r.story.strip().upper() for r in self.displacements if r.story}
        if weight and stories:
            return {s: weight / len(stories) for s in stories}
        return {}


def _is_weight_case(name: str) -> bool:
    normalized = name.upper().replace(" ", "")
    return any(k in normalized for k in ("WI", "W1", "WEIGHT", "POIDS", "DEAD", "G+Q", "DL"))


def _matches_case(load_case: str, base: str) -> bool:
    normalized = load_case.upper().replace(" ", "")
    return (
        normalized == base
        or normalized.startswith(base)
        or f"-{base}" in normalized
        or f"_{base}" in normalized
    )


def read_etabs(path: str | Path, joint_name: str = "") -> tuple[StructuralModel, EtabsResults, list[TabularTable]]:
    tables = read_tabular(path)
    name = joint_name.strip() or Path(path).stem
    model = build_model(tables, name)
    results = parse_etabs_results(tables)
    model.warnings.extend(results.warnings)
    return model, results, tables


def parse_etabs_results(tables: Sequence[TabularTable]) -> EtabsResults:
    units = _detect_units(tables)
    results = EtabsResults()

    for table in _find_tables(tables, "Modal Participating Mass Ratios", "Modal Participation Mass Ratios"):
        for row in table.iter_rows():
            mode = row.number("Mode", "ModeNumber", "Mode No")
            if mode is None:
                continue
            period = row.number("Period", "T") or 0.0
            results.modal.append(ModalRatio(
                int(mode), period,
                row.percent("UX", "Ux") or 0.0,
                row.percent("UY", "Uy") or 0.0,
                row.percent("UZ", "Uz") or 0.0,
                row.percent("SumUX", "Sum Ux") or 0.0,
                row.percent("SumUY", "Sum Uy") or 0.0,
                row.percent("SumUZ", "Sum Uz") or 0.0,
                row.percent("RX", "Rx") or 0.0,
                row.percent("RY", "Ry") or 0.0,
                row.percent("RZ", "Rz") or 0.0,
            ))
    if results.modal:
        maximum = max((m.sum_ux for m in results.modal), default=0.0)
        scale = 100.0 if 0.0 < maximum <= 1.5 else 1.0
        if scale > 1.0:
            for m in results.modal:
                m.ux *= scale; m.uy *= scale; m.uz *= scale
                m.sum_ux *= scale; m.sum_uy *= scale; m.sum_uz *= scale
                m.rx *= scale; m.ry *= scale; m.rz *= scale
            results.warnings.append("Participations massiques converties de fraction en pourcentage.")
        results.modal.sort(key=lambda m: m.mode)

    for table in _find_tables(tables, "Response Spectrum Base Reactions", "Base Reactions", "Reactions"):
        if table.resolve("FX") is None and table.resolve("F1") is None:
            continue
        for row in table.iter_rows():
            load_case = row.value("LoadCase", "Load Case", "Load Pattern", "Case", "Output Case")
            if not load_case:
                continue
            results.base_reactions.append(BaseReaction(
                load_case,
                units.to_kn(row.number("FX", "F1", "GlobalFX")),
                units.to_kn(row.number("FY", "F2", "GlobalFY")),
                units.to_kn(row.number("FZ", "F3", "GlobalFZ")),
                row.value("Story", "Level") or "",
            ))

    for table in _find_tables(tables, "Story Shears", "Story Forces"):
        if table.resolve("VX") is None and table.resolve("Vx") is None:
            continue
        for row in table.iter_rows():
            story = row.value("Story", "Level")
            load_case = row.value("LoadCase", "Load Case", "Load Case/Combo", "Case", "Output Case")
            if not story or not load_case:
                continue
            results.story_shears.append(StoryShear(
                story,
                load_case,
                row.value("Location", "Loc") or "",
                units.to_kn(row.number("VX", "Vx", "F1")),
                units.to_kn(row.number("VY", "Vy", "F2")),
                units.to_kn(row.number("P", "PX", "FZ")),
            ))

    for table in _find_tables(
        tables, "Diaphragm CM Displacements", "Diaphragm Center of Mass Displacements",
        "Story Displacements", "Joint Displacements",
    ):
        for row in table.iter_rows():
            story = row.value("Story", "Level")
            load_case = row.value("LoadCase", "Load Case", "Load Case/Combo", "Case", "Output Case")
            if not story or not load_case:
                continue
            results.displacements.append(Displacement(
                story,
                load_case,
                units.to_meters(row.number("UX", "Ux", "U1")),
                units.to_meters(row.number("UY", "Uy", "U2")),
                units.to_meters(row.number("UZ", "Uz", "U3")),
                row.value("Diaphragm", "Diaph") or "",
            ))

    for table in _find_tables(tables, "Column Forces", "Element Forces - Columns", "Element Forces Columns"):
        results.column_forces.extend(_read_forces(table, units, "Column"))
    for table in _find_tables(tables, "Beam Forces", "Element Forces - Beams", "Element Forces Beams"):
        results.beam_forces.extend(_read_forces(table, units, "Beam"))
    for table in _find_tables(tables, "Pier Forces", "Wall Forces", "Element Forces - Walls"):
        results.pier_forces.extend(_read_forces(table, units, "Pier"))

    for table in _find_tables(tables, "Centers Of Mass And Rigidity", "Center of Mass and Rigidity"):
        for row in table.iter_rows():
            story = row.value("Story", "Level")
            xcm = row.number("XCM", "X CM", "CenterMassX")
            ycm = row.number("YCM", "Y CM", "CenterMassY")
            if not story or xcm is None or ycm is None:
                continue
            results.centers.append(MassRigidityCenter(
                story,
                row.value("Diaphragm", "Diaph") or "",
                units.to_meters(xcm),
                units.to_meters(ycm),
                units.to_meters(row.number("XCR", "X CR", "CenterRigidityX") or xcm),
                units.to_meters(row.number("YCR", "Y CR", "CenterRigidityY") or ycm),
                row.number("MassX", "Mass X") or 0.0,
                row.number("MassY", "Mass Y") or 0.0,
            ))

    for table in _find_tables(tables, "Soil Pressure", "Soil Pressures", "Slab Soil Pressure", "Area Soil Pressure"):
        for row in table.iter_rows():
            label = row.value("Area", "UniqueName", "Label", "Joint", "Point", "Story")
            pressure = row.number("Pressure", "Q", "SoilPressure", "P", "PressureMax", "q")
            if not label or pressure is None:
                continue
            results.soil_pressures.append(SoilPressure(
                label,
                row.value("LoadCase", "Load Case", "Case", "Combo", "Output Case") or "",
                units.to_kpa(abs(pressure)),
                row.value("Story", "Level") or "",
                units.to_meters(row.number("Displacement", "UZ", "U3", "Settlement")),
                units.to_meters(row.number("Area", "ContactArea", "A")) if row.number("Area", "ContactArea", "A") else 0.0,
            ))

    if not results.modal:
        results.warnings.append("Table 'Modal Participating Mass Ratios' absente ou vide.")
    if not results.story_shears:
        results.warnings.append("Table 'Story Shears' absente ou vide : poids sismique non extrait.")
    if not results.displacements:
        results.warnings.append("Table de deplacements absente ou vide.")
    if not results.column_forces:
        results.warnings.append("Table 'Column Forces' absente ou vide.")
    if not results.centers:
        results.warnings.append("Table 'Centers Of Mass And Rigidity' absente ou vide.")
    return results


def _read_forces(table: TabularTable, units: UnitSystem, label: str) -> list[ForceResult]:
    rows: list[ForceResult] = []
    id_aliases = {
        "Column": ("Column", "UniqueName", "Frame", "Label"),
        "Beam": ("Beam", "UniqueName", "Frame", "Label"),
        "Pier": ("Pier", "Wall", "UniqueName", "Label"),
    }[label]
    for row in table.iter_rows():
        element = row.value(*id_aliases)
        load_case = row.value("LoadCase", "Load Case", "Load Case/Combo", "Case", "Output Case")
        if not element or not load_case:
            continue
        rows.append(ForceResult(
            element,
            load_case,
            row.value("Story", "Level") or "",
            units.to_meters(row.number("Station", "Loc", "Location")),
            units.to_kn(row.number("P", "P1")),
            units.to_kn(row.number("V2")),
            units.to_kn(row.number("V3")),
            units.to_knm(row.number("M2")),
            units.to_knm(row.number("M3")),
            units.to_meters(row.number("Length", "L")) if row.number("Length", "L") else 0.0,
        ))
    return rows


# ---------------------------------------------------------------------------
# Resultats SAFE
# ---------------------------------------------------------------------------

@dataclass
class SafeResults:
    soil_pressures: list[SoilPressure] = field(default_factory=list)
    footings: list[FootingResult] = field(default_factory=list)
    punching: list[dict[str, float | str]] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def read_safe(path: str | Path) -> SafeResults:
    tables = read_tabular(path)
    units = _detect_units(tables)
    result = SafeResults()

    for table in _find_tables(tables, "Soil Pressure", "Soil Pressures", "Slab Soil Pressure", "Area Soil Pressure"):
        for row in table.iter_rows():
            label = row.value("Area", "UniqueName", "Label", "Joint", "Point", "Story")
            pressure = row.number("Pressure", "Q", "SoilPressure", "P", "q")
            if not label or pressure is None:
                continue
            result.soil_pressures.append(SoilPressure(
                label,
                row.value("LoadCase", "Load Case", "Case", "Combo", "Output Case") or "",
                units.to_kpa(pressure),
                row.value("Story", "Level") or "",
                units.to_meters(row.number("Displacement", "UZ", "U3", "Settlement")),
            ))

    for table in _find_tables(tables, "Foundation Reactions", "Footing Forces", "Footing Design", "Isolated Footings"):
        for row in table.iter_rows():
            label = row.value("Footing", "Label", "Support", "Joint", "UniqueName")
            if not label:
                continue
            b = row.number("B", "Width", "B1")
            length = row.number("L", "Length", "B2")
            thickness = row.number("Thickness", "Depth", "H")
            result.footings.append(FootingResult(
                label,
                units.to_meters(b),
                units.to_meters(length),
                units.to_meters(thickness),
                units.to_kn(row.number("N", "P", "FZ", "Axial")),
                units.to_knm(row.number("MX", "Mx", "M1")),
                units.to_knm(row.number("MY", "My", "M2")),
                units.to_kpa(row.number("Qadm", "QAdm", "Allowable")),
                units.to_kpa(row.number("Qmax", "qmax")),
                units.to_kpa(row.number("Qmin", "qmin")),
                units.to_meters(row.number("Settlement", "Delta", "UZ")),
            ))

    for table in _find_tables(tables, "Punching Shear", "Slab Punching", "Punching"):
        for row in table.iter_rows():
            label = row.value("Area", "Label", "Point", "UniqueName")
            if not label:
                continue
            result.punching.append({
                "label": label,
                "vu": units.to_kn(row.number("Vu", "V", "Shear")),
                "vrd": units.to_kn(row.number("VRd", "Vrd", "Capacity")),
                "ratio": row.number("Ratio", "Utilization") or 0.0,
            })

    if not result.soil_pressures:
        result.warnings.append(
            "Aucune pression de sol lue depuis SAFE : les contraintes sont calculees "
            "a partir des reactions ETABS et de la portance saisie."
        )
    return result


# ---------------------------------------------------------------------------
# Configuration du projet
# ---------------------------------------------------------------------------

WILAYAS = [
    "01 - Adrar", "02 - Chlef", "03 - Laghouat", "04 - Oum El Bouaghi", "05 - Batna",
    "06 - Bejaia", "07 - Biskra", "08 - Bechar", "09 - Blida", "10 - Bouira",
    "11 - Tamanrasset", "12 - Tebessa", "13 - Tlemcen", "14 - Tiaret", "15 - Tizi Ouzou",
    "16 - Alger", "17 - Djelfa", "18 - Jijel", "19 - Setif", "20 - Saida",
    "21 - Skikda", "22 - Sidi Bel Abbes", "23 - Annaba", "24 - Guelma", "25 - Constantine",
    "26 - Medea", "27 - Mostaganem", "28 - M'Sila", "29 - Mascara", "30 - Ouargla",
    "31 - Oran", "32 - El Bayadh", "33 - Illizi", "34 - Bordj Bou Arreridj", "35 - Boumerdes",
    "36 - El Tarf", "37 - Tindouf", "38 - Tissemsilt", "39 - El Oued", "40 - Khenchela",
    "41 - Souk Ahras", "42 - Tipaza", "43 - Mila", "44 - Ain Defla", "45 - Naama",
    "46 - Ain Temouchent", "47 - Ghardaia", "48 - Relizane", "49 - Timimoun",
    "50 - Bordj Badji Mokhtar", "51 - Ouled Djellal", "52 - Beni Abbes", "53 - In Salah",
    "54 - In Guezzam", "55 - Touggourt", "56 - Djanet", "57 - El M'Ghair", "58 - El Meniaa",
]

OCCUPANCY_CASES = {
    "1 - Habitation / bureaux": 0.20,
    "2a - Public debout": 0.30,
    "2b - Public assis": 0.40,
    "3 - Entrepots / hangars": 0.50,
    "4 - Archives / reservoirs": 1.00,
    "5 - Autres locaux": 0.60,
}

BRACING_SYSTEMS = {
    "1 - Ossature BA (R=5,5)": (1, 5.5),
    "2 - Mixte equivalent ossature (R=5,5)": (2, 5.5),
    "3 - Ossature avec remplissage rigide (R=3,5)": (3, 3.5),
    "4 - Mixte equivalent voiles (R=4,5)": (4, 4.5),
    "5 - Voiles porteurs (R=4,5)": (5, 4.5),
    "6 - Ossature a noyau (R=3,0)": (6, 3.0),
    "7 - Console verticale (R=3,0)": (7, 3.0),
    "8 - Pendule inverse (R=2,0)": (8, 2.0),
    "9 - Voiles grandes dimensions (R=1,5)": (9, 1.5),
}

SEISMIC_ZONES = {
    "0 - Tres faible (A=0)": 0.0,
    "I - Faible (A=0,07)": 0.07,
    "II - Faible a moyenne (A=0,10)": 0.10,
    "III - Moyenne (A=0,15)": 0.15,
    "IV - Moyenne a elevee (A=0,20)": 0.20,
    "V - Elevee (A=0,25)": 0.25,
    "VI - Elevee (A=0,30)": 0.30,
}

SITE_CLASSES = {
    "S1 - Rocheux": "S1",
    "S2 - Ferme": "S2",
    "S3 - Meuble": "S3",
    "S4 - Tres meuble": "S4",
    "SS - Etude specifique": "SS",
}

IMPORTANCE_GROUPS = {
    "1A - Vitale (I=1,40)": ("1A", 1.40, 0.39),
    "1B - Grande importance (I=1,20)": ("1B", 1.20, 0.46),
    "2 - Importance moyenne (I=1,00)": ("2", 1.00, 0.55),
    "3 - Faible importance (I=0,80)": ("3", 0.80, 0.69),
}


@dataclass
class ProjectConfig:
    joint_name: str = "Joint A"
    project_name: str = "Projet"
    wilaya: str = "16 - Alger"
    zone: float = 0.15
    zone_label: str = "III - Moyenne (A=0,15)"
    site_class: str = "S3"
    importance_code: str = "2"
    importance_factor: float = 1.00
    nu_a: float = 0.55
    bracing_id: int = 1
    behavior_factor: float = 5.5
    occupancy: str = "1 - Habitation / bureaux"
    psi: float = 0.20
    plan_regular: bool = True
    elevation_regular: bool = True
    redundant_walls: bool = True
    qf_override: float | None = None
    ct_override: float | None = None
    damping_override: float | None = None
    lambda_override: float | None = None
    fc28: float = 25.0
    fe: float = 500.0
    gamma_b: float = 1.5
    gamma_s: float = 1.15
    concrete_density: float = 25.0
    steel_cover: float = 4.0  # cm
    drift_limit_fragile: float = 0.0050
    drift_limit_ductile: float = 0.0075
    pdelta_negative: float = 0.10
    pdelta_unstable: float = 0.20
    q_adm: float = 200.0  # kPa
    subgrade_modulus: float = 30000.0  # kN/m3 (Winkler)
    footing_concrete: float = 25.0
    footing_steel: float = 500.0
    include_penthouse: bool = True
    story_height_override: float | None = None
    manual_seismic_weight: float | None = None
    seismic_joint_provided: float | None = None
    settlement_limit_mm: float = 25.0
    notes: str = ""

    def zone_type(self) -> str:
        return "TYPE_1" if self.zone_label.strip().startswith(("IV", "V", "VI")) else "TYPE_2"


# ---------------------------------------------------------------------------
# Projet complet
# ---------------------------------------------------------------------------

@dataclass
class ParsedProject:
    model: StructuralModel
    results: EtabsResults
    config: ProjectConfig
    tables: list[TabularTable] = field(default_factory=list)
    safe: SafeResults | None = None
    source_files: list[str] = field(default_factory=list)
    safe_files: list[str] = field(default_factory=list)

    @property
    def warnings(self) -> list[str]:
        warnings = list(self.model.warnings) + list(self.results.warnings)
        if self.safe:
            warnings += self.safe.warnings
        return _dedupe(warnings)


def read_project(
    etabs_path: str | Path,
    config: ProjectConfig | None = None,
    safe_path: str | Path | None = None,
) -> ParsedProject:
    project_config = config or ProjectConfig()
    model, results, tables = read_etabs(etabs_path, project_config.joint_name)
    safe = read_safe(safe_path) if safe_path else None
    if safe and safe.soil_pressures:
        results.soil_pressures.extend(safe.soil_pressures)
    return ParsedProject(
        model=model,
        results=results,
        config=project_config,
        tables=tables,
        safe=safe,
        source_files=[str(Path(etabs_path))],
        safe_files=[str(Path(safe_path))] if safe_path else [],
    )


def _dedupe(values: Iterable[str]) -> list[str]:
    seen: set[str] = set()
    result = []
    for value in values:
        if value and value not in seen:
            seen.add(value)
            result.append(value)
    return result
