"""Tests automatiques de l'application Python Archilab RPA2024.

Execution :
    python -m unittest discover -s tests -v
"""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from parsers import (
    ProjectConfig,
    UnitSystem,
    read_project,
    read_tabular,
)
from rpa_engine import BaelSection, BaelSolver, Status, assess, evaluate_modal
from report_generator import generate_reports
from samples import make_demo_data


ROOT = Path(__file__).resolve().parents[1]
DEMO_ETABS = ROOT / "samples" / "DEMO_ETABS.txt"
DEMO_SAFE = ROOT / "samples" / "DEMO_SAFE.txt"


def demo_config(**overrides) -> ProjectConfig:
    config = ProjectConfig(
        project_name="Résidence Test",
        joint_name="Joint Test",
        wilaya="16 - Alger",
        zone=0.10,
        zone_label="II - Faible a moyenne (A=0,10)",
        site_class="S3",
        importance_code="2",
        importance_factor=1.0,
        nu_a=0.55,
        bracing_id=1,
        behavior_factor=5.5,
        fc28=25.0,
        fe=500.0,
        q_adm=200.0,
        subgrade_modulus=30000.0,
        seismic_joint_provided=250.0,
    )
    for key, value in overrides.items():
        setattr(config, key, value)
    return config


class UnitTests(unittest.TestCase):
    def test_etabs_units(self):
        units = UnitSystem.from_etabs("kN, m, C")
        self.assertAlmostEqual(units.to_kn(12.0), 12.0)
        self.assertAlmostEqual(units.to_meters(3.5), 3.5)
        self.assertAlmostEqual(units.to_kpa(180.0), 180.0)

    def test_tonf_cm_units(self):
        units = UnitSystem.from_etabs("Tonf, cm, C")
        self.assertAlmostEqual(units.to_kn(1.0), 9.80665, places=4)
        self.assertAlmostEqual(units.to_meters(100.0), 1.0, places=6)
        self.assertAlmostEqual(units.to_kpa(1.0), 98066.5, places=1)


class ParserTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.project = read_project(DEMO_ETABS, demo_config(), DEMO_SAFE)

    def test_model_counts(self):
        model = self.project.model
        self.assertEqual(model.num_stories, 5)
        self.assertEqual(len(model.columns), 80)
        self.assertEqual(len(model.beams), 120)
        self.assertEqual(len(model.walls), 20)
        self.assertEqual(len(model.slabs), 45)
        self.assertAlmostEqual(model.total_height, 15.0)
        self.assertEqual(model.num_bays_x, 3)
        self.assertEqual(model.num_bays_y, 3)

    def test_results(self):
        results = self.project.results
        self.assertEqual(len(results.modal), 12)
        self.assertEqual(len(results.centers), 5)
        self.assertEqual(len(results.soil_pressures), 4)
        self.assertAlmostEqual(results.base_shear("X"), 300.0)
        self.assertAlmostEqual(results.seismic_weight or 0, 4000.0)

    def test_xlsx_equivalent(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "demo.xlsx"
            make_demo_data.build_xlsx(make_demo_data.build_etabs_text(), path)
            tables = read_tabular(path)
            self.assertGreaterEqual(len(tables), 10)
            project = read_project(path, demo_config())
            self.assertEqual(project.model.num_stories, 5)
            self.assertEqual(len(project.model.columns), 80)
            self.assertAlmostEqual(project.results.base_shear("X") or 0, 300.0)

    def test_access_mdb_equivalent(self):
        mdb = ROOT / "samples" / "DEMO_ETABS.mdb"
        if not mdb.exists():
            self.skipTest("fixture Access absente")
        tables = read_tabular(mdb)
        self.assertGreaterEqual(len(tables), 10)
        project = read_project(mdb, demo_config())
        self.assertEqual(project.model.num_stories, 5)
        self.assertEqual(len(project.model.columns), 80)
        self.assertAlmostEqual(project.results.base_shear("X") or 0, 300.0)

    def test_model_is_not_hardcoded(self):
        old_grid, old_levels = make_demo_data.GRID, make_demo_data.LEVELS
        try:
            make_demo_data.GRID = [0.0, 5.0, 10.0, 15.0, 20.0]
            make_demo_data.LEVELS = 3
            text = make_demo_data.build_etabs_text()
            with tempfile.TemporaryDirectory() as temp:
                path = Path(temp) / "variante.txt"
                path.write_text(text, encoding="utf-8")
                project = read_project(path, demo_config())
                self.assertEqual(project.model.num_stories, 3)
                self.assertEqual(project.model.num_bays_x, 4)
                self.assertEqual(len(project.model.columns), 3 * 5 * 5)
                assessment = assess(project)
                self.assertGreaterEqual(len(assessment.chapters), 13)
        finally:
            make_demo_data.GRID = old_grid
            make_demo_data.LEVELS = old_levels


class BaelTests(unittest.TestCase):
    def test_simple_bending(self):
        solver = BaelSolver(25.0, 500.0)
        result = solver.solve(BaelSection(30.0, 35.0, 4.0, 4.0, 0.0, 95.0))
        self.assertGreater(result.required_cm2, 0)
        self.assertLess(result.required_cm2, 30.0)
        self.assertGreater(result.lever_arm, 0.2)

    def test_combined_bending(self):
        solver = BaelSolver(25.0, 500.0)
        result = solver.solve(BaelSection(50.0, 50.0, 4.0, 4.0, 350.0, 95.0, accidental=True))
        self.assertGreater(result.required_cm2, 0)
        self.assertIn(result.state, {
            "Section partiellement comprimée", "Section entièrement comprimée",
            "Section entièrement tendue", "Flexion simple",
        })


class EngineTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.project = read_project(DEMO_ETABS, demo_config(), DEMO_SAFE)
        cls.assessment = assess(cls.project)

    def test_status_and_chapters(self):
        assessment = self.assessment
        self.assertEqual(assessment.status, Status.OK)
        self.assertEqual(len(assessment.chapters), 13)
        self.assertEqual([c.number for c in assessment.chapters], list(range(1, 14)))
        self.assertEqual(assessment.failed_count, 0)

    def test_seismic(self):
        seismic = self.assessment.seismic
        self.assertEqual(seismic.resolved.zone_code, "II")
        self.assertAlmostEqual(seismic.resolved.qf, 1.0)
        self.assertAlmostEqual(seismic.resolved.qf, self.assessment.regularity.qf)
        self.assertEqual(seismic.modal.status, Status.OK)
        self.assertEqual(seismic.modal.modes_for_90, 8)
        for shear in seismic.base_shear:
            self.assertEqual(shear.status, Status.OK)
            self.assertGreaterEqual((shear.ratio or 0) * 100, 80)
        for drift in seismic.all_drifts:
            if drift.status != Status.NA:
                self.assertLessEqual(drift.reduced, drift.governing + 1e-9)
        for pdelta in seismic.all_pdelta:
            if pdelta.status != Status.NA:
                self.assertLessEqual(pdelta.theta, 0.10 + 1e-9)

    def test_normal_force_and_reinforcement(self):
        assessment = self.assessment
        self.assertTrue(assessment.seismic.columns)
        self.assertTrue(all(c.nu <= 0.35 for c in assessment.seismic.columns
                            if c.status != Status.NA))
        self.assertEqual(assessment.reinforcement.status, Status.OK)
        self.assertIn("C50x50", assessment.reinforcement.column_designs)
        self.assertTrue(assessment.reinforcement.column_designs["C50x50"]["label"].startswith("8T"))
        self.assertEqual(len(assessment.reinforcement.walls), 4)

    def test_foundations_and_dqe(self):
        assessment = self.assessment
        self.assertEqual(assessment.foundations.status, Status.OK)
        self.assertTrue(assessment.foundations.footings)
        footing = assessment.foundations.footings[0]
        self.assertLessEqual(footing.q_max, assessment.config.q_adm + 1e-6)
        self.assertGreaterEqual(footing.q_min, -1e-6)
        self.assertIsNotNone(assessment.dqe)
        self.assertGreater(assessment.dqe.total_concrete, 0)
        self.assertGreater(assessment.dqe.total_steel, 0)

    def test_strong_column_and_joint(self):
        joints = self.assessment.joints
        self.assertEqual(joints.status, Status.OK)
        self.assertGreater(len(joints.strong_column), 0)
        self.assertTrue(all(s.ratio >= 1.25 - 1e-9 for s in joints.strong_column))
        self.assertAlmostEqual(joints.seismic_joint.required_mm,
                               max(40.0, joints.seismic_joint.required_mm))

    def test_negative_foundation_case(self):
        project = read_project(DEMO_ETABS, demo_config(q_adm=50.0), DEMO_SAFE)
        assessment = assess(project)
        self.assertEqual(assessment.foundations.status, Status.FAILED)
        self.assertEqual(assessment.status, Status.FAILED)
        failed = [c for c in assessment.foundations.checks if c.status == Status.FAILED]
        self.assertTrue(failed)
        self.assertTrue(any("Portance" in c.label or "SAFE" in c.label for c in failed))

    def test_height_limit_warning(self):
        config = demo_config(zone=0.15, zone_label="III - Moyenne (A=0,15)")
        project = read_project(DEMO_ETABS, config)
        assessment = assess(project)
        self.assertIsNotNone(assessment.seismic.resolved.height_compliance)
        self.assertFalse(assessment.seismic.resolved.height_compliance[0])
        self.assertTrue(any("Limite de hauteur" in w for w in assessment.warnings))

    def test_modal_evaluation_direct(self):
        modal = evaluate_modal(self.project.results, 5)
        self.assertTrue(modal.check_a)
        self.assertEqual(modal.modes_for_90, 8)


class ReportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.project = read_project(DEMO_ETABS, demo_config(), DEMO_SAFE)
        cls.assessment = assess(cls.project)

    def test_docx_pdf_plans(self):
        with tempfile.TemporaryDirectory() as temp:
            result = generate_reports(self.assessment, temp, "NOTE_TEST")
            self.assertTrue(result.docx and result.docx.exists())
            self.assertTrue(result.pdf and result.pdf.exists())
            self.assertGreater(result.docx.stat().st_size, 20_000)
            self.assertGreater(result.pdf.stat().st_size, 20_000)
            self.assertTrue(result.plans)
            for path in result.plans.files:
                self.assertTrue(path.exists(), path)
                self.assertGreater(path.stat().st_size, 100, path)

            from docx import Document
            document = Document(result.docx)
            all_text = [p.text for p in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    for cell in row.cells:
                        all_text.append(cell.text)
            tags = [t for t in all_text if t.startswith("[FIGURE_CAPTURE_LOGICIEL")]
            self.assertGreaterEqual(len(tags), 13)
            cv = [t for t in all_text if t.strip().endswith("CV.")]
            self.assertGreaterEqual(len(cv), 13)
            headings = [p.text for p in document.paragraphs if p.style.name == "Heading 1"]
            self.assertGreaterEqual(len(headings), 14)  # sommaire + synthese + 13 chapitres

            try:
                from pypdf import PdfReader
            except ImportError:
                PdfReader = None
            if PdfReader:
                reader = PdfReader(result.pdf)
                self.assertGreaterEqual(len(reader.pages), 20)

    def test_dxf_audit(self):
        try:
            import ezdxf
        except ImportError:
            self.skipTest("ezdxf non installé")
        with tempfile.TemporaryDirectory() as temp:
            result = generate_reports(self.assessment, temp, "NOTE_DXF", include_plans=True)
            self.assertTrue(result.plans)
            for path in (result.plans.formwork_dxf, result.plans.reinforcement_dxf):
                self.assertIsNotNone(path)
                document = ezdxf.readfile(path)
                auditor = document.audit()
                self.assertEqual(len(auditor.errors), 0)
                layers = {layer.dxf.name for layer in document.layers}
                for required in ("AXES", "COLONNES", "POUTRES", "VOILES", "PLANCHERS",
                                 "COTES", "TEXTE", "CADRES", "BARRES", "HACHURES", "CARTOUCHE"):
                    self.assertIn(required, layers)
                self.assertGreater(len(document.modelspace()), 50)


if __name__ == "__main__":
    unittest.main()
